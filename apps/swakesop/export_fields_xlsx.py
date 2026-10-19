"""
Run this once on the server to generate the category fields spreadsheet:
  cd ~/docker/netinventory_v2
  pip3 install openpyxl --quiet
  python3 export_fields_xlsx.py
Then SCP the output file to your machine.
"""
import re, sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from seed_category_fields_v1 import SCHEMA
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

CATEGORY_LABELS = {
    "universal": "Universal (All Assets)",
    "servers": "Servers",
    "raspberry_pi": "Raspberry Pi",
    "network_equipment": "Network Equipment",
    "cameras": "Cameras",
    "power": "Power (UPS-PDU)",
    "storage": "Storage (NAS-DAS)",
    "laptops": "Laptops",
    "lxc_vms": "LXC - VMs",
    "cables_adapters": "Cables & Adapters",
    "iot_devices": "IoT Devices",
    "media_gaming": "Media & Gaming",
    "mobile_tablets": "Mobile & Tablets",
    "pcs_workstations": "PCs & Workstations",
    "peripherals": "Peripherals",
}

GROUP_LABELS = {
    "core": "Core", "hardware_specs": "Hardware Specs", "storage": "Storage",
    "os_software": "OS / Software", "support": "Support", "board": "Board",
    "os": "OS", "hardware_add_ons": "Hardware Add-ons", "device_specs": "Device Specs",
    "management": "Management", "power": "Power", "optics_sensor": "Optics / Sensor",
    "connectivity": "Connectivity", "recording": "Recording",
    "ups_pdu_specs": "UPS / PDU Specs", "load_monitoring": "Load Monitoring",
    "battery": "Battery", "array_enclosure": "Array / Enclosure",
    "drives_installed": "Drives Installed", "hardware": "Hardware",
    "support_assignment": "Support / Assignment", "virtual_machine_config": "VM Config",
    "services": "Services", "cable_adapter_details": "Cable / Adapter Details",
    "condition_storage": "Condition / Storage", "device_details": "Device Details",
    "placement": "Placement", "automation": "Automation", "gaming": "Gaming",
    "smart_streaming": "Smart / Streaming", "monitor_specific": "Monitor-specific",
    "printer_specific": "Printer-specific", "mobile": "Mobile",
}

HDR_BG = "1E2C4A"; CAT_BG = "2563EB"; ROW_ODD = "0F1D3D"; ROW_EVEN = "16223F"
WHITE = "F2F6FF"; MUTED = "94A3B8"; GOLD = "FACC15"

def fill(c): return PatternFill("solid", fgColor=c)
def bdr():
    s = Side(style="thin", color="334155")
    return Border(left=s, right=s, top=s, bottom=s)
def safe_title(t): return re.sub(r'[\\/*?:\[\]]', '-', t)[:31]

wb = Workbook()
ws = wb.active
ws.title = "All Fields"

for col, h in enumerate(["Category", "Section", "Field Label", "Field Type", "Example / Placeholder"], 1):
    c = ws.cell(row=1, column=col, value=h)
    c.font = Font(bold=True, color=WHITE, size=11)
    c.fill = fill(HDR_BG); c.alignment = Alignment(horizontal="center", vertical="center"); c.border = bdr()
ws.row_dimensions[1].height = 22

row_num = 2
for cat_slug, groups in SCHEMA.items():
    cat_label = CATEGORY_LABELS.get(cat_slug, cat_slug.replace("_", " ").title())
    for grp_key, fields in groups.items():
        grp_label = GROUP_LABELS.get(grp_key, grp_key.replace("_", " ").title())
        for field_key, label, field_type, placeholder in fields:
            ws.append([cat_label, grp_label, label, field_type, placeholder or ""])
            f = fill(ROW_ODD if row_num % 2 == 0 else ROW_EVEN)
            for col in range(1, 6):
                c = ws.cell(row=row_num, column=col)
                c.fill = f; c.font = Font(color=WHITE, size=10)
                c.alignment = Alignment(vertical="center", wrap_text=True); c.border = bdr()
            ws.cell(row=row_num, column=1).font = Font(color="7DD3FC", size=10)
            ws.cell(row=row_num, column=2).font = Font(color=MUTED, size=10)
            ws.cell(row=row_num, column=4).font = Font(color=GOLD, size=10)
            row_num += 1

ws.column_dimensions["A"].width = 26; ws.column_dimensions["B"].width = 22
ws.column_dimensions["C"].width = 28; ws.column_dimensions["D"].width = 14
ws.column_dimensions["E"].width = 45
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:E{row_num-1}"

for cat_slug, groups in SCHEMA.items():
    cat_label = CATEGORY_LABELS.get(cat_slug, cat_slug.replace("_", " ").title())
    ws2 = wb.create_sheet(title=safe_title(cat_label))
    ws2.merge_cells("A1:E1")
    t = ws2["A1"]
    t.value = cat_label; t.font = Font(bold=True, color=WHITE, size=13)
    t.fill = fill(CAT_BG); t.alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[1].height = 28
    for col, h in enumerate(["Section", "Field Label", "Field Type", "Example / Placeholder", "Your Notes"], 1):
        c = ws2.cell(row=2, column=col, value=h)
        c.font = Font(bold=True, color=WHITE, size=10)
        c.fill = fill(HDR_BG); c.alignment = Alignment(horizontal="center", vertical="center"); c.border = bdr()
    ws2.row_dimensions[2].height = 20
    r = 3
    for grp_key, fields in groups.items():
        grp_label = GROUP_LABELS.get(grp_key, grp_key.replace("_", " ").title())
        for field_key, label, field_type, placeholder in fields:
            ws2.append([grp_label, label, field_type, placeholder or "", ""])
            f = fill(ROW_ODD if r % 2 == 0 else ROW_EVEN)
            for col in range(1, 6):
                c = ws2.cell(row=r, column=col)
                c.fill = f; c.font = Font(color=WHITE, size=10)
                c.alignment = Alignment(vertical="center", wrap_text=True); c.border = bdr()
            ws2.cell(row=r, column=1).font = Font(color=MUTED, size=10)
            ws2.cell(row=r, column=3).font = Font(color=GOLD, size=10)
            r += 1
    ws2.column_dimensions["A"].width = 22; ws2.column_dimensions["B"].width = 28
    ws2.column_dimensions["C"].width = 14; ws2.column_dimensions["D"].width = 45
    ws2.column_dimensions["E"].width = 30
    ws2.freeze_panes = "A3"

out = Path(__file__).parent / "swops_category_fields.xlsx"
wb.save(out)
print(f"Saved: {out}")
